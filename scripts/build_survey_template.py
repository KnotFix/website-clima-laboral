"""
Arma la plantilla gratuita de encuesta de clima laboral, en español y en inglés.

    python scripts/build_survey_template.py

Requiere `openpyxl`. En Windows con Excel instalado (y `pywin32` + `pymupdf`)
ademas hace tres cosas que openpyxl no puede:

1. Abre cada archivo en Excel, recalcula y **verifica que ninguna formula de la
   hoja de resultados de error** con los datos de ejemplo. Si alguna lo da, el
   script falla.
2. Lo guarda desde Excel, para que el archivo lleve los valores ya calculados:
   una vista previa que no recalcula (la del correo, la del telefono) muestra
   los resultados en vez de celdas vacias.
3. Exporta la parte de arriba de la hoja de resultados como PNG: es la captura
   que muestra la pagina del sitio.

Despues de guardar desde Excel se **limpian los metadatos** del archivo: Excel
escribe el usuario de Windows como autor y la ruta local de la carpeta, y eso
no puede salir en un archivo publico.

**El texto sale de `src/content/resources/survey_template.json`**, el mismo que
dibuja la pagina. Las preguntas se escribieron desde cero para este recurso: no
son ni derivan de las del producto, y la hoja de resultados hace solo cuentas
simples (porcentaje favorable y eNPS estandar). Ver la regla en el encabezado
de `src/content/resources/nav.js`.
"""

import json
import math
import random
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
CONTENT = ROOT / "src" / "content" / "resources" / "survey_template.json"
PUBLIC = ROOT / "public"

ROWS = 300  # personas que entran en la hoja de respuestas
EXAMPLE_ROWS = 20
FIRST_DATA_ROW = 3  # fila 1: dimensiones; fila 2: P1..P30
LAST_DATA_ROW = FIRST_DATA_ROW + ROWS - 1
QUESTIONS_PER_DIMENSION = 5

# Paleta: tinta y blanco, el morado del sitio solo como acento.
INK = "18181B"
MUTED = "6B6B76"
ACCENT = "5B21B6"
BAND = "F2F2F4"
EXAMPLE_GREY = "D9D9DE"
INPUT_FILL = "FFF7D6"
LINE = "C9C9D1"

TITLE_FONT = Font(name="Calibri", size=18, bold=True, color=INK)
SUBTITLE_FONT = Font(name="Calibri", size=11, color=MUTED)
SECTION_FONT = Font(name="Calibri", size=13, bold=True, color=ACCENT)
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BOLD = Font(name="Calibri", size=11, bold=True, color=INK)
BODY = Font(name="Calibri", size=11, color=INK)
NOTE = Font(name="Calibri", size=10, italic=True, color=MUTED)
LINK = Font(name="Calibri", size=12, bold=True, color=ACCENT, underline="single")

HEADER_FILL = PatternFill("solid", fgColor=INK)
BAND_FILL = PatternFill("solid", fgColor=BAND)
INPUT = PatternFill("solid", fgColor=INPUT_FILL)
THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNDERLINE = Border(bottom=THIN)

WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

LINE_HEIGHT = 15  # puntos por renglon de Calibri 11


def height_for(text, width, font_size=11):
    """Alto de fila para que `text` entre en una celda de `width` unidades.

    openpyxl no ajusta el alto de las filas, y una celda combinada tampoco la
    ajusta Excel: se estima por caracteres por renglon.
    """
    chars_per_line = max(1, int(width * 1.15 * 11 / font_size))
    lines = sum(
        max(1, math.ceil(len(paragraph) / chars_per_line))
        for paragraph in str(text).split("\n")
    )
    return lines * LINE_HEIGHT * font_size / 11 + 4


def all_questions(content):
    return [q for dimension in content["dimensions"] for q in dimension["questions"]]


def build_guide(ws, content):
    sheet = content["guide_sheet"]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 90

    row = 2
    ws.cell(row, 2, content["file_title"]).font = TITLE_FONT
    ws.row_dimensions[row].height = 28
    row += 1
    ws.cell(row, 2, sheet["subtitle"]).font = SUBTITLE_FONT
    row += 2

    def section(title):
        nonlocal row
        ws.cell(row, 2, title).font = SECTION_FONT
        ws.row_dimensions[row].height = 22
        row += 1

    def paragraph(text, font=BODY, column=3, width=90):
        nonlocal row
        cell = ws.cell(row, column, text)
        cell.font = font
        cell.alignment = WRAP
        ws.row_dimensions[row].height = height_for(text, width)
        row += 1

    section(sheet["steps_title"])
    for index, step in enumerate(content["steps"], start=1):
        number = ws.cell(row, 2, index)
        number.font = BOLD
        number.alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 3, step["title"]).font = BOLD
        row += 1
        paragraph(step["body"])
    row += 1

    section(sheet["scale_title"])
    for item in content["scale"]:
        value = ws.cell(row, 2, item["value"])
        value.font = BOLD
        value.alignment = Alignment(horizontal="center")
        value.border = BOX
        label = ws.cell(row, 3, item["label"])
        label.font = BODY
        label.border = BOX
        if item["value"] >= 3:
            value.fill = BAND_FILL
            label.fill = BAND_FILL
        row += 1
    paragraph(content["favorable_note"], NOTE)
    paragraph(sheet["scale_note"], NOTE)
    row += 1

    section(sheet["dimensions_title"])
    for index, dimension in enumerate(content["dimensions"]):
        first = index * QUESTIONS_PER_DIMENSION + 1
        last = first + QUESTIONS_PER_DIMENSION - 1
        label = sheet["dimension_range"].format(
            name=dimension["name"], first=first, last=last, prefix=content["question_prefix"]
        )
        ws.cell(row, 3, label).font = BODY
        row += 1
    paragraph(sheet["dimensions_note"], NOTE)
    row += 1

    section(content["anonymity_title"])
    for tip in content["anonymity_tips"]:
        ws.cell(row, 2, "•").alignment = Alignment(horizontal="center", vertical="top")
        paragraph(tip)
    row += 1

    section(content["limits_title"])
    for limit in content["limits"]:
        ws.cell(row, 2, "•").alignment = Alignment(horizontal="center", vertical="top")
        paragraph(limit)
    row += 1
    closing = ws.cell(row, 2, content["limits_closing"])
    closing.font = LINK
    closing.hyperlink = "https://censuma.com"
    ws.row_dimensions[row].height = 22

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "censuma.com"


def build_survey(ws, content):
    sheet = content["survey_sheet"]
    ws.sheet_view.showGridLines = False
    widths = {"A": 6, "B": 62, "C": 11, "D": 11, "E": 11, "F": 15}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    full_width = sum(widths.values())

    ws.cell(1, 1, content["file_title"]).font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 28

    intro = ws.cell(2, 1, sheet["intro"])
    intro.font = BODY
    intro.alignment = WRAP
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = height_for(sheet["intro"], full_width)

    ws.cell(3, 1, sheet["area_label"]).font = BOLD
    ws.merge_cells("A3:B3")
    for column in range(3, 7):
        ws.cell(3, column).border = UNDERLINE
    ws.row_dimensions[3].height = 24

    header_row = 5
    headers = [sheet["number_header"], sheet["question_header"]] + [
        f"{item['value']}\n{item['label']}" for item in content["scale"]
    ]
    for column, text in enumerate(headers, start=1):
        cell = ws.cell(header_row, column, text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER if column != 2 else WRAP_CENTER
        cell.border = BOX
    ws.row_dimensions[header_row].height = 32

    row = header_row + 1
    number = 1
    for dimension in content["dimensions"]:
        cell = ws.cell(row, 1, dimension["name"])
        cell.font = BOLD
        cell.fill = BAND_FILL
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.row_dimensions[row].height = 20
        row += 1
        for question in dimension["questions"]:
            ws.cell(row, 1, number).alignment = Alignment(horizontal="center", vertical="center")
            text = ws.cell(row, 2, question)
            text.font = BODY
            text.alignment = WRAP_CENTER
            for column in range(1, 7):
                ws.cell(row, column).border = BOX
            ws.row_dimensions[row].height = max(24, height_for(question, widths["B"]))
            row += 1
            number += 1

    row += 1
    ws.cell(row, 1, sheet["enps_title"]).font = SECTION_FONT
    row += 1
    enps = ws.cell(row, 1, content["enps"]["question"])
    enps.font = BODY
    enps.alignment = WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.row_dimensions[row].height = height_for(content["enps"]["question"], full_width)
    row += 1
    scale = ws.cell(row, 1, "     ".join(str(n) for n in range(11)))
    scale.font = Font(name="Calibri", size=14, bold=True, color=INK)
    scale.alignment = CENTER
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.row_dimensions[row].height = 28
    row += 1
    ends = ws.cell(row, 1, f"{content['enps']['low']}          {content['enps']['high']}")
    ends.font = NOTE
    ends.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    row += 2

    ws.cell(row, 1, sheet["open_title"]).font = SECTION_FONT
    row += 1
    for question in content["open_questions"]:
        cell = ws.cell(row, 1, question)
        cell.font = BOLD
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        row += 1
        for _ in range(4):
            for column in range(1, 7):
                ws.cell(row, column).border = UNDERLINE
            ws.row_dimensions[row].height = 24
            row += 1
        row += 1

    ws.cell(row, 1, sheet["thanks"]).font = BOLD

    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.oddFooter.center.text = "censuma.com"
    ws.print_area = f"A1:F{row}"


def build_responses(ws, content, rng):
    sheet = content["responses_sheet"]
    questions = all_questions(content)
    enps_col = 3 + len(questions)
    open_cols = [enps_col + 1 + i for i in range(len(content["open_questions"]))]
    last_col = open_cols[-1]

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 26
    for index in range(len(questions)):
        ws.column_dimensions[get_column_letter(3 + index)].width = 5.5
    ws.column_dimensions[get_column_letter(enps_col)].width = 7
    for column in open_cols:
        ws.column_dimensions[get_column_letter(column)].width = 42

    # Fila 1: el nombre de cada dimension sobre sus cinco columnas.
    for column in (1, 2):
        ws.cell(1, column).fill = HEADER_FILL
    for index, dimension in enumerate(content["dimensions"]):
        start = 3 + index * QUESTIONS_PER_DIMENSION
        end = start + QUESTIONS_PER_DIMENSION - 1
        cell = ws.cell(1, start, dimension["name"])
        cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL if index % 2 == 0 else PatternFill("solid", fgColor="3A3A42")
        cell.alignment = CENTER
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
    for column in range(enps_col, last_col + 1):
        ws.cell(1, column).fill = HEADER_FILL
    ws.row_dimensions[1].height = 30

    # Fila 2: los encabezados de cada columna. Cada P lleva su pregunta como
    # nota, para no tener que ir a la otra hoja a ver que era la P17.
    labels = [sheet["number_header"], sheet["area_header"]]
    labels += [f"{content['question_prefix']}{i}" for i in range(1, len(questions) + 1)]
    labels += [content["enps"]["label"]]
    labels += [f"{sheet['open_header_prefix']} {i}" for i in range(1, len(open_cols) + 1)]
    for column, text in enumerate(labels, start=1):
        cell = ws.cell(2, column, text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BOX
    for index, question in enumerate(questions):
        note = Comment(question, "Censuma")
        note.width = 260
        note.height = 70
        ws.cell(2, 3 + index).comment = note
    ws.cell(2, enps_col).comment = Comment(content["enps"]["question"], "Censuma", 90, 260)
    for i, column in enumerate(open_cols):
        ws.cell(2, column).comment = Comment(content["open_questions"][i], "Censuma", 60, 260)
    ws.row_dimensions[2].height = 22

    for offset in range(ROWS):
        row = FIRST_DATA_ROW + offset
        ws.cell(row, 1, offset + 1).alignment = Alignment(horizontal="center")
        ws.cell(row, 1).font = Font(name="Calibri", size=10, color=MUTED)

    # Filas de ejemplo. Cada dimension tiene su propio sesgo para que los
    # resultados de ejemplo no salgan todos iguales.
    dimension_bias = [0.15, 0.0, 0.35, -0.05, -0.35, 0.2]
    open_examples = sheet["example_comments"]
    for offset in range(EXAMPLE_ROWS):
        row = FIRST_DATA_ROW + offset
        area = sheet["example_areas"][offset % len(sheet["example_areas"])]
        ws.cell(row, 2, f"{sheet['example_area_prefix']}: {area}")
        person_bias = rng.gauss(0, 0.45)
        for index in range(len(questions)):
            bias = dimension_bias[index // QUESTIONS_PER_DIMENSION] + person_bias
            value = round(2.75 + bias + rng.gauss(0, 0.7))
            ws.cell(row, 3 + index, min(4, max(1, value)))
        enps = round(7 + person_bias * 3 + rng.gauss(0, 1.6))
        ws.cell(row, enps_col, min(10, max(0, enps)))
        if offset % 3 == 0:
            pair = open_examples[(offset // 3) % len(open_examples)]
            ws.cell(row, open_cols[0], pair[0])
            ws.cell(row, open_cols[1], pair[1])

    scale_range = f"C{FIRST_DATA_ROW}:{get_column_letter(2 + len(questions))}{LAST_DATA_ROW}"
    enps_letter = get_column_letter(enps_col)
    enps_range = f"{enps_letter}{FIRST_DATA_ROW}:{enps_letter}{LAST_DATA_ROW}"

    scale_validation = DataValidation(
        type="whole",
        operator="between",
        formula1="1",
        formula2="4",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle=sheet["validation_title_scale"],
        error=sheet["validation_error_scale"],
        showInputMessage=True,
        prompt=sheet["validation_prompt_scale"],
    )
    scale_validation.add(scale_range)
    enps_validation = DataValidation(
        type="whole",
        operator="between",
        formula1="0",
        formula2="10",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle=sheet["validation_title_enps"],
        error=sheet["validation_error_enps"],
        showInputMessage=True,
        prompt=sheet["validation_prompt_enps"],
    )
    enps_validation.add(enps_range)
    ws.add_data_validation(scale_validation)
    ws.add_data_validation(enps_validation)

    # Gris mientras la fila sea de ejemplo: al borrarla con Suprimir el gris se
    # va solo, y una fila cargada con datos reales nunca sale gris.
    prefix = sheet["example_area_prefix"]
    data_range = f"A{FIRST_DATA_ROW}:{get_column_letter(last_col)}{LAST_DATA_ROW}"
    ws.conditional_formatting.add(
        data_range,
        FormulaRule(
            formula=[f'LEFT($B{FIRST_DATA_ROW},{len(prefix)})="{prefix}"'],
            fill=PatternFill("solid", fgColor=EXAMPLE_GREY, bgColor=EXAMPLE_GREY),
            font=Font(color="55555E"),
        ),
    )

    for column in range(3, enps_col + 1):
        for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
            ws.cell(row, column).alignment = Alignment(horizontal="center")

    ws.freeze_panes = f"C{FIRST_DATA_ROW}"
    return {
        "first_letter": "C",
        "last_question_letter": get_column_letter(2 + len(questions)),
        "enps_letter": enps_letter,
    }


def build_results(ws, content, refs):
    sheet = content["results_sheet"]
    responses = content["sheets"]["responses"]
    data = f"'{responses}'!"
    first, last_q = refs["first_letter"], refs["last_question_letter"]

    def col_range(start_letter, end_letter=None):
        end_letter = end_letter or start_letter
        return f"{data}${start_letter}${FIRST_DATA_ROW}:${end_letter}${LAST_DATA_ROW}"

    all_answers = col_range(first, last_q)
    enps_range = col_range(refs["enps_letter"])

    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 66
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 13

    ws.cell(1, 1, sheet["title"]).font = TITLE_FONT
    ws.row_dimensions[1].height = 28
    subtitle = ws.cell(2, 1, sheet["subtitle"])
    subtitle.font = SUBTITLE_FONT
    subtitle.alignment = WRAP
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = height_for(sheet["subtitle"], 100, 11)

    def section(row, title):
        ws.cell(row, 1, title).font = SECTION_FONT
        ws.row_dimensions[row].height = 22

    def metric(row, label, formula, number_format, bold=False):
        ws.cell(row, 2, label).font = BOLD if bold else BODY
        value = ws.cell(row, 3, formula)
        value.number_format = number_format
        value.font = Font(name="Calibri", size=13 if bold else 11, bold=True, color=INK)
        value.alignment = Alignment(horizontal="right")
        return value

    # Participacion
    section(4, sheet["participation_title"])
    invited = metric(5, sheet["invited_label"], 25, "0")
    invited.fill = INPUT
    invited.border = BOX
    # Cuenta las filas que tienen al menos una respuesta (en las 30 preguntas o
    # en el eNPS): MMULT suma cada fila y SUMPRODUCT cuenta las que dan > 0.
    width = len(all_questions(content)) + 1
    answered = (
        f"=SUMPRODUCT(--(MMULT(--({data}${first}${FIRST_DATA_ROW}:${refs['enps_letter']}${LAST_DATA_ROW}<>\"\"),"
        f"ROW($A$1:$A${width})^0)>0))"
    )
    metric(6, sheet["responses_label"], answered, "0")
    metric(7, sheet["participation_label"], "=IFERROR(C6/C5,\"\")", "0%", bold=True)

    # Resultado general
    section(9, sheet["general_title"])
    metric(
        10,
        sheet["general_label"],
        f"=IFERROR(COUNTIF({all_answers},\">=3\")/COUNT({all_answers}),\"\")",
        "0%",
        bold=True,
    )
    note = ws.cell(11, 2, sheet["general_note"])
    note.font = NOTE
    note.alignment = WRAP

    # eNPS
    section(13, sheet["enps_title"])
    metric(14, sheet["promoters_label"], f"=IFERROR(COUNTIF({enps_range},\">=9\")/COUNT({enps_range}),\"\")", "0%")
    metric(
        15,
        sheet["passives_label"],
        f"=IFERROR(COUNTIFS({enps_range},\">=7\",{enps_range},\"<=8\")/COUNT({enps_range}),\"\")",
        "0%",
    )
    metric(16, sheet["detractors_label"], f"=IFERROR(COUNTIF({enps_range},\"<=6\")/COUNT({enps_range}),\"\")", "0%")
    metric(17, sheet["enps_label"], "=IFERROR(ROUND((C14-C16)*100,0),\"\")", "+0;-0;0", bold=True)
    enps_note = ws.cell(18, 2, sheet["enps_note"])
    enps_note.font = NOTE
    enps_note.alignment = WRAP

    def table_header(row, first_label):
        labels = ["", first_label, sheet["favorable_header"], sheet["answers_header"]]
        for column, text in enumerate(labels, start=1):
            cell = ws.cell(row, column, text)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = CENTER if column != 2 else WRAP_CENTER
        ws.row_dimensions[row].height = 22

    # Por dimension
    section(20, sheet["dimension_title"])
    table_header(21, sheet["dimension_header"])
    row = 22
    dimension_first = row
    for index, dimension in enumerate(content["dimensions"]):
        start = get_column_letter(3 + index * QUESTIONS_PER_DIMENSION)
        end = get_column_letter(3 + index * QUESTIONS_PER_DIMENSION + QUESTIONS_PER_DIMENSION - 1)
        block = col_range(start, end)
        ws.cell(row, 1, index + 1).alignment = Alignment(horizontal="center")
        ws.cell(row, 2, dimension["name"]).font = BODY
        favorable = ws.cell(row, 3, f"=IFERROR(COUNTIF({block},\">=3\")/COUNT({block}),\"\")")
        favorable.number_format = "0%"
        favorable.font = BOLD
        count = ws.cell(row, 4, f"=COUNT({block})")
        count.number_format = "0"
        count.font = Font(name="Calibri", size=10, color=MUTED)
        for column in range(1, 5):
            ws.cell(row, column).border = Border(bottom=THIN)
        row += 1
    dimension_last = row - 1
    screenshot_last_row = dimension_last

    # Por pregunta
    row += 1
    section(row, sheet["question_title"])
    row += 1
    table_header(row, sheet["question_header"])
    row += 1
    question_first = row
    for index, question in enumerate(all_questions(content)):
        letter = get_column_letter(3 + index)
        column = col_range(letter)
        ws.cell(row, 1, f"{content['question_prefix']}{index + 1}").alignment = Alignment(horizontal="center", vertical="center")
        text = ws.cell(row, 2, question)
        text.font = BODY
        text.alignment = WRAP_CENTER
        favorable = ws.cell(row, 3, f"=IFERROR(COUNTIF({column},\">=3\")/COUNT({column}),\"\")")
        favorable.number_format = "0%"
        favorable.font = BOLD
        favorable.alignment = Alignment(vertical="center")
        count = ws.cell(row, 4, f"=COUNT({column})")
        count.font = Font(name="Calibri", size=10, color=MUTED)
        count.alignment = Alignment(vertical="center")
        for c in range(1, 5):
            ws.cell(row, c).border = Border(bottom=THIN)
        ws.row_dimensions[row].height = max(20, height_for(question, 66))
        row += 1
    question_last = row - 1

    bar = DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="A78BFA")
    ws.conditional_formatting.add(f"C{dimension_first}:C{dimension_last}", bar)
    ws.conditional_formatting.add(
        f"C{question_first}:C{question_last}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="A78BFA"),
    )

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "censuma.com"
    ws.print_area = f"A1:D{question_last}"
    return {
        "formula_range": f"C5:D{question_last}",
        "screenshot_range": f"$A$1:$D${screenshot_last_row}",
    }


def build(lang, content):
    # La misma semilla en los dos idiomas: los ejemplos (y la captura) dan igual.
    rng = random.Random("censuma-plantilla")
    wb = Workbook()
    guide = wb.active
    guide.title = content["sheets"]["guide"]
    survey = wb.create_sheet(content["sheets"]["survey"])
    responses = wb.create_sheet(content["sheets"]["responses"])
    results = wb.create_sheet(content["sheets"]["results"])

    build_guide(guide, content)
    build_survey(survey, content)
    refs = build_responses(responses, content, rng)
    results_info = build_results(results, content, refs)

    wb.properties.title = content["file_title"]
    wb.properties.creator = "Censuma"
    wb.properties.lastModifiedBy = "Censuma"
    wb.properties.subject = content["file_title"]
    wb.properties.keywords = "censuma.com"
    wb.active = 0

    target = PUBLIC / content["download_path"].lstrip("/")
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target, results_info


# ---------------------------------------------------------------------------
# Excel: recalculo, verificacion, captura y limpieza de metadatos
# ---------------------------------------------------------------------------

ERROR_CODES = {-2146826281, -2146826246, -2146826259, -2146826288, -2146826252, -2146826265, -2146826273}


def finalize_with_excel(path, content, results_info):
    import pymupdf
    import win32com.client

    sheets = content["sheets"]
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(path))
        excel.CalculateFull()

        # 1. Ninguna formula de la hoja de resultados da error.
        results = wb.Worksheets(sheets["results"])
        errors = []
        formulas = 0
        for cell in results.Range(results_info["formula_range"]):
            if cell.HasFormula:
                formulas += 1
                value = cell.Value
                if isinstance(value, int) and value in ERROR_CODES:
                    errors.append(f"{cell.Address}: {cell.Text}")
                if cell.Text.startswith("#"):
                    errors.append(f"{cell.Address}: {cell.Text}")
        if errors:
            raise SystemExit(f"{path.name}: formulas con error: {errors}")

        summary = {
            "formulas": formulas,
            "respuestas": results.Range("C6").Value,
            "participacion": results.Range("C7").Text,
            "favorable_general": results.Range("C10").Text,
            "enps": results.Range("C17").Text,
            "dimensiones": [results.Range(f"C{r}").Text for r in range(22, 28)],
        }

        # 2. Guardar desde Excel, con los valores calculados.
        wb.Worksheets(1).Activate()
        wb.Save()

        # 3. La captura: la parte de arriba de Resultados, hasta la tabla por
        # dimension. Se cambia el area de impresion solo para exportar y se
        # cierra sin guardar.
        results.PageSetup.PrintArea = results_info["screenshot_range"]
        results.PageSetup.Zoom = False
        results.PageSetup.FitToPagesWide = 1
        results.PageSetup.FitToPagesTall = 1
        results.PageSetup.CenterFooter = ""
        pdf = Path(tempfile.gettempdir()) / f"{path.stem}-resultados.pdf"
        results.ExportAsFixedFormat(0, str(pdf))
        wb.Close(SaveChanges=False)
    finally:
        excel.Quit()

    image_target = PUBLIC / content["preview_image"].lstrip("/")
    document = pymupdf.open(pdf)
    page = document[0]
    pix = page.get_pixmap(dpi=200)
    pix.save(image_target)
    document.close()
    pdf.unlink(missing_ok=True)
    _trim_image(image_target)

    _scrub_metadata(path)
    return summary


def _trim_image(path, margin=48):
    from PIL import Image, ImageChops

    image = Image.open(path).convert("RGB")
    background = Image.new("RGB", image.size, (255, 255, 255))
    box = ImageChops.difference(image, background).getbbox()
    if box:
        left, top, right, bottom = box
        box = (
            max(0, left - margin),
            max(0, top - margin),
            min(image.width, right + margin),
            min(image.height, bottom + margin),
        )
        image = image.crop(box)
    image.save(path, optimize=True)


def _scrub_metadata(path):
    """Saca del .xlsx lo que Excel agrega de la maquina que lo guardo."""
    replacements = {
        "docProps/core.xml": [
            (r"<dc:creator>.*?</dc:creator>", "<dc:creator>Censuma</dc:creator>"),
            (r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy>Censuma</cp:lastModifiedBy>"),
            (r"<cp:lastPrinted>.*?</cp:lastPrinted>", ""),
        ],
        "docProps/app.xml": [
            (r"<Company>.*?</Company>", "<Company>Censuma</Company>"),
            (r"<Manager>.*?</Manager>", ""),
        ],
        "xl/workbook.xml": [
            (r"<mc:AlternateContent[^>]*>\s*<mc:Choice[^>]*>\s*<x15ac:absPath[^>]*/>\s*</mc:Choice>\s*</mc:AlternateContent>", ""),
            (r"<x15ac:absPath[^>]*/>", ""),
        ],
    }
    temp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in replacements:
                text = data.decode("utf-8")
                for pattern, replacement in replacements[item.filename]:
                    text = re.sub(pattern, replacement, text, flags=re.S)
                data = text.encode("utf-8")
            target.writestr(item, data)
    shutil.move(temp, path)


def main():
    content = json.loads(CONTENT.read_text(encoding="utf-8"))
    use_excel = sys.platform == "win32" and "--no-excel" not in sys.argv
    for lang in ("es", "en"):
        path, results_info = build(lang, content[lang])
        print(f"{lang}: {path.relative_to(ROOT)}")
        if use_excel:
            summary = finalize_with_excel(path, content[lang], results_info)
            print(f"    verificado en Excel: {summary}")


if __name__ == "__main__":
    main()
